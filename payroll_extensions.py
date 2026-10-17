"""Optional payroll catalog, exceptions, revisions, audit, and documents.

All features default off except Excel bonuses (already live). Disabled flags
leave calculate_monthly_payroll math unchanged.
"""
from __future__ import annotations

import html as html_lib
import io
import json
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Sequence, Tuple

from payroll_policy_schema import DEFAULT_CURRENCY, DEFAULT_DECIMAL_PLACES, money, tenant_ids

FEATURE_FLAGS = (
    "enable_named_components",
    "enable_tax",
    "enable_loans",
    "enable_advances",
    "enable_provident",
    "enable_insurance",
    "enable_bonuses",
    "enable_individual_adjustments",
    "enable_salary_revisions",
)

DEFAULT_FEATURE_FLAGS = {
    "enable_named_components": 0,
    "enable_tax": 0,
    "enable_loans": 0,
    "enable_advances": 0,
    "enable_provident": 0,
    "enable_insurance": 0,
    "enable_bonuses": 1,
    "enable_individual_adjustments": 0,
    "enable_salary_revisions": 0,
}

VALID_COMPONENT_KINDS = ("earning", "deduction")
VALID_COMPONENT_CATEGORIES = (
    "allowance", "bonus", "tax", "provident", "insurance", "loan", "advance", "other",
)
VALID_CALC_MODES = ("fixed", "percent_basic", "percent_gross", "per_working_day", "per_hour")
VALID_SCOPES = ("org", "branch", "shift", "department", "designation", "employee")
SCOPE_RANK = {
    "employee": 0,
    "designation": 1,
    "department": 2,
    "shift": 3,
    "branch": 4,
    "org": 5,
}
# percent_gross uses basic + non-percent-gross named earnings (no circular reference).
VALID_ADJ_TYPES = ("rest_day", "bonus", "earning", "deduction", "overtime", "arrears")
VALID_ADJ_MODES = ("multiplier", "fixed")

CATEGORY_FLAG = {
    "tax": "enable_tax",
    "loan": "enable_loans",
    "advance": "enable_advances",
    "provident": "enable_provident",
    "insurance": "enable_insurance",
    "bonus": "enable_bonuses",
}


def flag_on(settings: Any, name: str) -> bool:
    if not settings:
        return bool(DEFAULT_FEATURE_FLAGS.get(name))
    data = dict(settings) if not isinstance(settings, dict) and hasattr(settings, "keys") else (settings or {})
    if name not in data or data.get(name) is None:
        return bool(DEFAULT_FEATURE_FLAGS.get(name))
    return bool(int(data.get(name) or 0))


def ensure_payroll_extension_tables(cursor) -> None:
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS payroll_component_types (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            organization_id TEXT NOT NULL DEFAULT 'DEFAULT_ORG',
            branch_id TEXT NOT NULL DEFAULT 'DEFAULT_BRANCH',
            name TEXT NOT NULL,
            code TEXT NOT NULL,
            kind TEXT NOT NULL CHECK(kind IN ('earning', 'deduction')),
            category TEXT NOT NULL DEFAULT 'allowance',
            calc_mode TEXT NOT NULL DEFAULT 'fixed',
            default_amount REAL NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(organization_id, branch_id, code)
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS payroll_employee_components (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            organization_id TEXT NOT NULL DEFAULT 'DEFAULT_ORG',
            branch_id TEXT NOT NULL DEFAULT 'DEFAULT_BRANCH',
            user_id INTEGER NOT NULL,
            component_type_id INTEGER NOT NULL,
            amount REAL NOT NULL DEFAULT 0,
            percent REAL NOT NULL DEFAULT 0,
            effective_from TEXT NOT NULL,
            effective_to TEXT,
            active INTEGER NOT NULL DEFAULT 1,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS salary_revisions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            organization_id TEXT NOT NULL DEFAULT 'DEFAULT_ORG',
            branch_id TEXT NOT NULL DEFAULT 'DEFAULT_BRANCH',
            user_id INTEGER NOT NULL,
            effective_from TEXT NOT NULL,
            basic_salary REAL NOT NULL DEFAULT 0,
            allowances REAL NOT NULL DEFAULT 0,
            deductions REAL NOT NULL DEFAULT 0,
            ot_rate REAL NOT NULL DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS payroll_adjustments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            organization_id TEXT NOT NULL DEFAULT 'DEFAULT_ORG',
            branch_id TEXT NOT NULL DEFAULT 'DEFAULT_BRANCH',
            user_id INTEGER NOT NULL,
            payroll_month TEXT NOT NULL,
            work_date TEXT,
            adj_type TEXT NOT NULL,
            pay_mode TEXT NOT NULL DEFAULT 'fixed',
            amount REAL NOT NULL DEFAULT 0,
            multiplier REAL NOT NULL DEFAULT 0,
            note TEXT NOT NULL DEFAULT '',
            source_month TEXT,
            created_by TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS payroll_duplicate_component_report (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            organization_id TEXT NOT NULL DEFAULT 'DEFAULT_ORG',
            branch_id TEXT NOT NULL DEFAULT 'DEFAULT_BRANCH',
            user_id INTEGER,
            component_type_id INTEGER NOT NULL,
            assignment_ids TEXT NOT NULL DEFAULT '',
            overlap_from TEXT,
            overlap_to TEXT,
            detected_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    _migrate_component_assignment_columns(cursor)
    _refresh_duplicate_component_report(cursor)
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS payroll_audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            organization_id TEXT NOT NULL DEFAULT 'DEFAULT_ORG',
            branch_id TEXT NOT NULL DEFAULT 'DEFAULT_BRANCH',
            payroll_month TEXT,
            action TEXT NOT NULL,
            actor TEXT,
            detail TEXT NOT NULL DEFAULT '{}',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
    )


def _columns(cursor, table: str) -> set:
    return {row[1] for row in cursor.execute(f"PRAGMA table_info({table})").fetchall()}


def _migrate_component_assignment_columns(cursor) -> None:
    """Idempotent. Rollback: DROP the added columns is not possible in SQLite; restore a backup."""
    type_cols = _columns(cursor, "payroll_component_types")
    for column, spec in (
        ("min_amount", "REAL"),
        ("max_amount", "REAL"),
        ("rounding", "TEXT NOT NULL DEFAULT 'cent'"),
    ):
        if column not in type_cols:
            cursor.execute(f"ALTER TABLE payroll_component_types ADD COLUMN {column} {spec}")
    asg_cols = _columns(cursor, "payroll_employee_components")
    for column, spec in (
        ("scope_type", "TEXT NOT NULL DEFAULT 'employee'"),
        ("scope_id", "TEXT"),
        ("excluded", "INTEGER NOT NULL DEFAULT 0"),
        ("min_amount", "REAL"),
        ("max_amount", "REAL"),
        ("rounding", "TEXT NOT NULL DEFAULT 'cent'"),
        ("calc_mode", "TEXT"),
    ):
        if column not in asg_cols:
            cursor.execute(f"ALTER TABLE payroll_employee_components ADD COLUMN {column} {spec}")
    cursor.execute(
        "UPDATE payroll_employee_components SET scope_type = 'employee' WHERE scope_type IS NULL OR scope_type = ''"
    )


def _dates_overlap(a_from: Optional[str], a_to: Optional[str], b_from: Optional[str], b_to: Optional[str]) -> bool:
    start_a = a_from or "0000-01-01"
    start_b = b_from or "0000-01-01"
    end_a = a_to or "9999-12-31"
    end_b = b_to or "9999-12-31"
    return start_a <= end_b and start_b <= end_a


def _refresh_duplicate_component_report(cursor) -> None:
    if not cursor.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='payroll_employee_components'"
    ).fetchone():
        return
    cursor.execute("DELETE FROM payroll_duplicate_component_report")
    fetched = cursor.execute(
        "SELECT * FROM payroll_employee_components WHERE active = 1"
    ).fetchall()
    columns = [item[0] for item in (cursor.description or [])]
    rows = []
    for record in fetched:
        if hasattr(record, "keys"):
            rows.append(dict(record))
        else:
            rows.append(dict(zip(columns, record)))
    seen = set()
    for i, left in enumerate(rows):
        for right in rows[i + 1 :]:
            if int(left.get("component_type_id") or 0) != int(right.get("component_type_id") or 0):
                continue
            if str(left.get("organization_id")) != str(right.get("organization_id")):
                continue
            if str(left.get("branch_id")) != str(right.get("branch_id")):
                continue
            if str(left.get("scope_type") or "employee") != str(right.get("scope_type") or "employee"):
                continue
            if int(left.get("user_id") or 0) != int(right.get("user_id") or 0):
                continue
            if str(left.get("scope_id") or "") != str(right.get("scope_id") or ""):
                continue
            if int(left.get("excluded") or 0) or int(right.get("excluded") or 0):
                continue
            if not _dates_overlap(left.get("effective_from"), left.get("effective_to"), right.get("effective_from"), right.get("effective_to")):
                continue
            key = tuple(sorted((left["id"], right["id"])))
            if key in seen:
                continue
            seen.add(key)
            cursor.execute(
                """INSERT INTO payroll_duplicate_component_report (
                       organization_id, branch_id, user_id, component_type_id, assignment_ids, overlap_from, overlap_to
                   ) VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (
                    left.get("organization_id"),
                    left.get("branch_id"),
                    left.get("user_id"),
                    left.get("component_type_id"),
                    f"{left['id']},{right['id']}",
                    max(left.get("effective_from") or "", right.get("effective_from") or ""),
                    min(left.get("effective_to") or "9999-12-31", right.get("effective_to") or "9999-12-31"),
                ),
            )


def assignment_matches_employee(row: Dict[str, Any], employee: Dict[str, Any], org: str, branch: str) -> bool:
    scope = str(row.get("scope_type") or "employee").strip().lower()
    sid = str(row.get("scope_id") or "").strip().lower()
    if scope == "employee":
        return int(row.get("user_id") or 0) == int(employee.get("employee_id") or employee.get("user_id") or 0)
    if scope == "designation":
        return sid == str(employee.get("designation") or "").strip().lower()
    if scope == "department":
        return sid == str(employee.get("department") or "").strip().lower()
    if scope == "shift":
        return str(row.get("scope_id") or "") == str(employee.get("shift_id") or "")
    if scope == "branch":
        return sid == str(branch).strip().lower() or str(row.get("scope_id") or "") == str(branch)
    if scope == "org":
        return True
    return False


def resolve_named_assignments(
    all_assignments: Sequence[Dict[str, Any]],
    employee: Dict[str, Any],
    org: str,
    branch: str,
    as_of: date,
) -> List[Dict[str, Any]]:
    as_of_s = as_of.isoformat()
    excluded_types = set()
    candidates: List[Dict[str, Any]] = []
    for row in all_assignments:
        if not row.get("active", 1):
            continue
        start = str(row.get("effective_from") or "0000-01-01")[:10]
        end = str(row.get("effective_to") or "9999-12-31")[:10]
        if as_of_s < start or as_of_s > end:
            continue
        if not assignment_matches_employee(row, employee, org, branch):
            continue
        if int(row.get("excluded") or 0):
            excluded_types.add(int(row.get("component_type_id") or 0))
            continue
        candidates.append(dict(row))
    by_type: Dict[int, Dict[str, Any]] = {}
    for row in candidates:
        type_id = int(row.get("component_type_id") or 0)
        if type_id in excluded_types and str(row.get("scope_type") or "employee") != "employee":
            continue
        if type_id in excluded_types and str(row.get("scope_type") or "employee") == "employee":
            continue
        rank = SCOPE_RANK.get(str(row.get("scope_type") or "employee"), 9)
        current = by_type.get(type_id)
        if current is None or rank < SCOPE_RANK.get(str(current.get("scope_type") or "employee"), 9):
            by_type[type_id] = row
    return list(by_type.values())


def write_audit(cursor, action: str, organization_id: Optional[str], branch_id: Optional[str],
                payroll_month: Optional[str] = None, actor: Optional[str] = None, detail: Optional[Dict[str, Any]] = None) -> None:
    org, branch = tenant_ids(organization_id, branch_id)
    cursor.execute(
        """INSERT INTO payroll_audit_log (
               organization_id, branch_id, payroll_month, action, actor, detail
           ) VALUES (?, ?, ?, ?, ?, ?)""",
        (org, branch, payroll_month, action, actor or "Admin", json.dumps(detail or {})),
    )


def _as_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    text = str(value)[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def package_as_of(
    revisions: Sequence[Dict[str, Any]],
    fallback: Dict[str, Any],
    as_of: date,
) -> Dict[str, float]:
    chosen = None
    for row in revisions:
        start = _as_date(row.get("effective_from"))
        if start and start <= as_of:
            if chosen is None or start >= (_as_date(chosen.get("effective_from")) or date.min):
                chosen = row
    src = chosen or fallback
    return {
        "basic_salary": float(src.get("basic_salary") or 0),
        "allowances": float(src.get("allowances") or 0),
        "deductions": float(src.get("deductions") or 0),
        "ot_rate": float(src.get("ot_rate") or 0),
    }


def prorate_packages(
    revisions: Sequence[Dict[str, Any]],
    fallback: Dict[str, Any],
    eligible_start: date,
    schedule_end: date,
    month_start: date,
    month_end: date,
    decimal_places: int,
    enabled: bool,
) -> Tuple[float, float, float, float]:
    """Return (basic, allowances, deductions, ot_rate) for the month slice."""
    if not enabled or not revisions:
        days_in_month = (month_end - month_start).days + 1
        active = max(0, (schedule_end - eligible_start).days + 1)
        factor = 1.0 if active >= days_in_month else (active / days_in_month if days_in_month else 1.0)
        return (
            money(float(fallback.get("basic_salary") or 0) * factor, decimal_places),
            money(float(fallback.get("allowances") or 0) * factor, decimal_places),
            money(float(fallback.get("deductions") or 0) * factor, decimal_places),
            money(float(fallback.get("ot_rate") or 0), decimal_places),
        )
    days_in_month = (month_end - month_start).days + 1 or 1
    basic = allow = deduct = 0.0
    current = eligible_start
    last_ot = float(fallback.get("ot_rate") or 0)
    while current <= schedule_end:
        pkg = package_as_of(revisions, fallback, current)
        basic += float(pkg["basic_salary"]) / days_in_month
        allow += float(pkg["allowances"]) / days_in_month
        deduct += float(pkg["deductions"]) / days_in_month
        last_ot = float(pkg["ot_rate"] or last_ot)
        current += timedelta(days=1)
    return money(basic, decimal_places), money(allow, decimal_places), money(deduct, decimal_places), money(last_ot, decimal_places)


def category_allowed(category: str, settings: Any) -> bool:
    flag = CATEGORY_FLAG.get(str(category or "").strip().lower())
    if not flag:
        return True
    return flag_on(settings, flag)


def named_component_totals(
    assignments: Sequence[Dict[str, Any]],
    types_by_id: Dict[int, Dict[str, Any]],
    basic_salary: float,
    as_of: date,
    settings: Any,
    decimal_places: int,
    employee: Optional[Dict[str, Any]] = None,
    organization_id: str = "DEFAULT_ORG",
    branch_id: str = "DEFAULT_BRANCH",
    working_days: float = 0.0,
    overtime_hours: float = 0.0,
) -> Tuple[float, float, List[Dict[str, Any]]]:
    resolved = list(assignments)
    if employee is not None:
        resolved = resolve_named_assignments(assignments, employee, organization_id, branch_id, as_of)
    prepared: List[Tuple[Dict[str, Any], Dict[str, Any], str]] = []
    for row in resolved:
        spec = types_by_id.get(int(row.get("component_type_id") or 0))
        if not spec or not spec.get("active", 1):
            continue
        if not category_allowed(spec.get("category") or "", settings):
            continue
        mode = str(row.get("calc_mode") or spec.get("calc_mode") or "fixed").strip().lower()
        prepared.append((row, spec, mode))

    def _amount_for(row: Dict[str, Any], spec: Dict[str, Any], mode: str, gross_base: float) -> float:
        if mode == "percent_basic":
            amount = basic_salary * float(row.get("percent") or spec.get("default_amount") or 0) / 100.0
        elif mode == "percent_gross":
            amount = gross_base * float(row.get("percent") or spec.get("default_amount") or 0) / 100.0
        elif mode == "per_working_day":
            amount = float(row.get("amount") or spec.get("default_amount") or 0) * float(working_days or 0)
        elif mode == "per_hour":
            amount = float(row.get("amount") or spec.get("default_amount") or 0) * float(overtime_hours or 0)
        else:
            amount = float(row.get("amount") if row.get("amount") not in (None, "") else spec.get("default_amount") or 0)
        lo = row.get("min_amount") if row.get("min_amount") not in (None, "") else spec.get("min_amount")
        hi = row.get("max_amount") if row.get("max_amount") not in (None, "") else spec.get("max_amount")
        if lo not in (None, ""):
            amount = max(amount, float(lo))
        if hi not in (None, "") and float(hi) > 0:
            amount = min(amount, float(hi))
        return money(max(0.0, amount), decimal_places)

    first_pass_earn = 0.0
    for row, spec, mode in prepared:
        if spec.get("kind") == "earning" and mode != "percent_gross":
            first_pass_earn += _amount_for(row, spec, mode, basic_salary)
    gross_base = basic_salary + first_pass_earn
    earnings = 0.0
    deductions = 0.0
    lines: List[Dict[str, Any]] = []
    for row, spec, mode in prepared:
        amount = _amount_for(row, spec, mode, gross_base)
        kind = spec.get("kind") or "earning"
        lines.append({
            "name": spec.get("name"),
            "code": spec.get("code"),
            "kind": kind,
            "category": spec.get("category"),
            "amount": amount,
            "scope_type": row.get("scope_type") or "employee",
        })
        if kind == "deduction":
            deductions += amount
        else:
            earnings += amount
    return money(earnings, decimal_places), money(deductions, decimal_places), lines


def apply_adjustments(
    adjustments: Sequence[Dict[str, Any]],
    daily_rate: float,
    off_day_dates: Sequence[str],
    decimal_places: int,
) -> Dict[str, Any]:
    extra_earn = 0.0
    extra_deduct = 0.0
    arrears = 0.0
    covered_off_days = set()
    lines: List[Dict[str, Any]] = []
    off_set = {str(d) for d in off_day_dates}
    for adj in adjustments:
        adj_type = str(adj.get("adj_type") or "").strip().lower()
        mode = str(adj.get("pay_mode") or "fixed").strip().lower()
        work_date = str(adj.get("work_date") or "").strip()
        if mode == "multiplier":
            amount = money(daily_rate * float(adj.get("multiplier") or 0), decimal_places)
        else:
            amount = money(float(adj.get("amount") or 0), decimal_places)
        if adj_type == "rest_day":
            covered_off_days.add(work_date or "")
            extra_earn += amount
        elif adj_type in ("bonus", "earning", "overtime"):
            extra_earn += amount
        elif adj_type == "deduction":
            extra_deduct += amount
        elif adj_type == "arrears":
            arrears += amount
            if amount >= 0:
                extra_earn += amount
            else:
                extra_deduct += abs(amount)
        lines.append({
            "type": adj_type,
            "date": work_date,
            "amount": amount,
            "note": adj.get("note") or "",
            "source_month": adj.get("source_month"),
        })
    remaining_off = [d for d in off_set if d not in covered_off_days]
    return {
        "extra_earn": money(extra_earn, decimal_places),
        "extra_deduct": money(extra_deduct, decimal_places),
        "arrears": money(arrears, decimal_places),
        "remaining_off_days": remaining_off,
        "lines": lines,
    }


def simple_pdf_bytes(title: str, lines: Sequence[str]) -> bytes:
    """Minimal one-page PDF without extra packages."""
    content_lines = ["BT /F1 14 Tf 50 780 Td 18 TL ({}) Tj T*".format(_pdf_escape(title))]
    content_lines.append("/F1 10 Tf 14 TL")
    for line in lines:
        content_lines.append("({}) Tj T*".format(_pdf_escape(line)))
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("latin-1", "replace")
    objs = []
    objs.append(b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n")
    objs.append(b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n")
    objs.append(
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
    )
    objs.append(b"4 0 obj << /Length %d >> stream\n" % len(stream) + stream + b"\nendstream endobj\n")
    objs.append(b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n")
    parts = [b"%PDF-1.4\n"]
    offsets = [0]
    for obj in objs:
        offsets.append(sum(len(p) for p in parts))
        parts.append(obj)
    xref_pos = sum(len(p) for p in parts)
    xref = [b"xref\n0 6\n0000000000 65535 f \n"]
    for off in offsets[1:]:
        xref.append("{:010d} 00000 n \n".format(off).encode("ascii"))
    parts.extend(xref)
    parts.append(b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref_pos)
    return b"".join(parts)


def _pdf_escape(text: str) -> str:
    return str(text or "").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")[:180]


def payroll_excel_bytes(payroll: Dict[str, Any], currency: str = DEFAULT_CURRENCY) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll"
    headers = [
        "Employee ID", "Name", "Department", "Basic", "Allowances", "Overtime",
        "Bonus", "Rest-day extra", "Adjustments", "Arrears", "Gross", "Deductions", "Net", "Status",
    ]
    ws.append(["Payroll", payroll.get("payroll_month"), payroll.get("status"), currency])
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F2942", end_color="0F2942", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    for row in payroll.get("records") or []:
        ws.append([
            row.get("employee_id"),
            row.get("employee_name"),
            row.get("department"),
            float(row.get("basic_salary") or 0),
            float(row.get("allowances") or 0),
            float(row.get("overtime_amount") or 0),
            float(row.get("bonus_amount") or 0),
            float(row.get("off_day_extra_pay") or 0),
            float(row.get("adjustment_amount") or 0),
            float(row.get("arrears_amount") or 0),
            float(row.get("gross_salary") or 0),
            float(row.get("deduction_amount") or 0),
            float(row.get("net_salary") or 0),
            row.get("status"),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _pdf_money(value: Any, currency: str, places: int) -> str:
    amount = float(value or 0)
    return "{} {:,.{}f}".format(currency, amount, places)


def _pdf_rows_from_payslip(payslip: Dict[str, Any]) -> Tuple[List[Tuple[str, str]], List[Tuple[str, str]]]:
    emp = payslip.get("employee") or {}
    currency = payslip.get("currency_code") or DEFAULT_CURRENCY
    places = int(payslip.get("decimal_places") or DEFAULT_DECIMAL_PLACES)
    deductions = payslip.get("deductions") or {}
    earnings: List[Tuple[str, str]] = [
        ("Basic pay", _pdf_money(emp.get("basic_salary"), currency, places)),
    ]
    pay_lines = emp.get("pay_lines") or []
    if isinstance(pay_lines, str):
        try:
            pay_lines = json.loads(pay_lines)
        except (TypeError, ValueError):
            pay_lines = []
    has_named_earn = any(
        str(item.get("kind") or "") == "earning" and float(item.get("amount") or 0)
        for item in (pay_lines or [])
    )
    if float(emp.get("allowances") or 0) and not has_named_earn:
        earnings.append(("Allowances", _pdf_money(emp.get("allowances"), currency, places)))
    if float(emp.get("overtime_amount") or 0):
        hours = emp.get("overtime_hours")
        label = "Overtime"
        if hours:
            label = "Overtime ({} hrs)".format(hours)
        earnings.append((label, _pdf_money(emp.get("overtime_amount"), currency, places)))
    if float(emp.get("bonus_amount") or 0) and not any(
        str(item.get("category") or "") == "bonus" for item in (pay_lines or [])
    ):
        earnings.append(("Bonus", _pdf_money(emp.get("bonus_amount"), currency, places)))
    if float(emp.get("off_day_extra_pay") or 0):
        earnings.append(("Rest-day extra", _pdf_money(emp.get("off_day_extra_pay"), currency, places)))
    adj_lines = [
        item for item in (pay_lines or [])
        if (item.get("note") is not None or item.get("source_month") is not None or (item.get("type") and not item.get("kind") and not item.get("category")))
    ]
    if adj_lines:
        for adj in adj_lines:
            amt = float(adj.get("amount") or 0)
            if not amt:
                continue
            adj_type = str(adj.get("type") or "adjustment").replace("_", " ").title()
            note = str(adj.get("note") or "").strip()
            date_str = str(adj.get("date") or "").strip()
            label = f"Adj: {adj_type}"
            if note:
                label += f" ({note[:16]})"
            elif date_str:
                label += f" ({date_str})"
            if amt >= 0:
                earnings.append((label, _pdf_money(amt, currency, places)))
    elif float(emp.get("adjustment_amount") or 0):
        earnings.append(("Adjustments", _pdf_money(emp.get("adjustment_amount"), currency, places)))
    if float(emp.get("arrears_amount") or 0):
        earnings.append(("Arrears", _pdf_money(emp.get("arrears_amount"), currency, places)))
    cuts: List[Tuple[str, str]] = []
    for item in pay_lines:
        amount = float(item.get("amount") or 0)
        if amount == 0:
            continue
        name = str(item.get("name") or item.get("type") or "Pay item")
        kind = str(item.get("kind") or item.get("category") or "")
        money_text = _pdf_money(amount, currency, places)
        if kind in ("deduction", "tax", "loan", "advance", "provident", "insurance"):
            cuts.append((name, money_text))
        elif kind in ("earning", "allowance", "bonus"):
            if name.lower() not in ("basic", "basic pay"):
                earnings.append((name, money_text))
    if float(deductions.get("absence") or 0):
        cuts.append(("Missed-day cut", _pdf_money(deductions.get("absence"), currency, places)))
    if float(deductions.get("unpaid_leave") or 0):
        cuts.append(("Unpaid leave", _pdf_money(deductions.get("unpaid_leave"), currency, places)))
    if float(deductions.get("half_leave") or 0):
        cuts.append(("Half-day leave cut", _pdf_money(deductions.get("half_leave"), currency, places)))
    if float(deductions.get("attendance_shortfall") or 0):
        cuts.append(("Days not present", _pdf_money(deductions.get("attendance_shortfall"), currency, places)))
    if float(deductions.get("other") or 0) and not any(row[0] for row in cuts if row[0] not in ("Missed-day cut", "Unpaid leave", "Half-day leave cut", "Days not present")):
        cuts.append(("Other cuts", _pdf_money(deductions.get("other"), currency, places)))
    if not cuts:
        cuts.append(("No cuts", _pdf_money(0, currency, places)))
    return earnings, cuts


def _payslip_letterhead(payslip: Dict[str, Any]) -> Dict[str, str]:
    lh = payslip.get("letterhead") or {}
    company = payslip.get("company") or {}
    defaults = {
        "payslip_company_name": company.get("name") or "QIntellect Technologies",
        "payslip_tagline": "DRIVING THE FUTURE TECHNOLOGY POWERED BY QUANTUM & AI",
        "payslip_phone": company.get("phone") or "+92 302 9633999",
        "payslip_email": company.get("email") or "info@qintellecttechnologies.com",
        "payslip_website": "qintellecttechnologies.com",
        "payslip_address": company.get("address") or company.get("branch") or "Main Branch, Rahim Yar Khan",
        "payslip_ntn": "J756903",
        "payslip_footer": "© 2026 QIntellect Technologies. All rights reserved.",
        "payslip_powered": "Powered by QIntellect Technologies",
        "payslip_bank_title": "QIntellect Technologies",
        "payslip_bank_name": "United Bank Limited (UBL)",
        "payslip_bank_iban": "PK46UNIL0109000400912116",
    }
    out = {}
    for key, fallback in defaults.items():
        text = str(lh.get(key) or fallback or "").strip()
        out[key] = text or str(fallback)
    return out


def _payslip_invoice_items(payslip: Dict[str, Any]) -> List[Dict[str, Any]]:
    emp = payslip.get("employee") or {}
    deductions = payslip.get("deductions") or {}
    items: List[Dict[str, Any]] = []

    def add(description: str, amount: Any, qty: Any = 1, force: bool = False) -> None:
        value = float(amount or 0)
        quantity = float(qty or 0) or 1.0
        if not force and value == 0:
            return
        items.append({
            "description": description,
            "qty": quantity,
            "unit": value / quantity if quantity else value,
            "amount": value,
        })

    add("Basic pay", emp.get("basic_salary"), 1, True)
    pay_lines = emp.get("pay_lines") or []
    if isinstance(pay_lines, str):
        try:
            pay_lines = json.loads(pay_lines)
        except (TypeError, ValueError):
            pay_lines = []
    named_earn = any(
        str(item.get("kind") or "") == "earning" and float(item.get("amount") or 0)
        for item in (pay_lines or [])
    )
    named_bonus = any(str(item.get("category") or "") == "bonus" for item in (pay_lines or []))
    if not named_earn:
        add("Allowances", emp.get("allowances"))
    ot_hours = float(emp.get("overtime_hours") or 0)
    ot_amount = float(emp.get("overtime_amount") or 0)
    if ot_amount:
        add("Overtime", ot_amount, ot_hours if ot_hours else 1)
    if not named_bonus:
        add("Bonus", emp.get("bonus_amount"))
    adj_lines = [
        item for item in (pay_lines or [])
        if (item.get("note") is not None or item.get("source_month") is not None or (item.get("type") and not item.get("kind") and not item.get("category")))
    ]
    if adj_lines:
        for adj in adj_lines:
            amt = float(adj.get("amount") or 0)
            if not amt:
                continue
            adj_type = str(adj.get("type") or "adjustment").replace("_", " ").title()
            note = str(adj.get("note") or "").strip()
            date_str = str(adj.get("date") or "").strip()
            desc = f"Adjustment: {adj_type}"
            if note:
                desc += f" ({note})"
            elif date_str:
                desc += f" ({date_str})"
            add(desc, abs(amt), 1, True)
    elif float(emp.get("adjustment_amount") or 0):
        add("Adjustments", emp.get("adjustment_amount"))
    add("Arrears", emp.get("arrears_amount"))
    for item in pay_lines or []:
        amount = float(item.get("amount") or 0)
        if not amount:
            continue
        name = str(item.get("name") or item.get("type") or "Pay item")
        kind = str(item.get("kind") or item.get("category") or "")
        if kind in ("deduction", "tax", "loan", "advance", "provident", "insurance"):
            add(name, amount, 1, True)
        elif kind in ("earning", "allowance", "bonus") and name.lower() not in ("basic", "basic pay"):
            add(name, amount, 1, True)
    add("Missed-day cut", deductions.get("absence"))
    add("Unpaid leave", deductions.get("unpaid_leave"))
    add("Half-day leave cut", deductions.get("half_leave"))
    add("Days not present", deductions.get("attendance_shortfall"))
    if not items:
        add("Basic pay", 0, 1, True)
    return items


def payslip_invoice_html(payslip: Dict[str, Any]) -> str:
    """Invoice-paper HTML shared by the on-screen payslip and print/PDF views."""
    emp = payslip.get("employee") or {}
    lh = _payslip_letterhead(payslip)
    currency = payslip.get("currency_code") or DEFAULT_CURRENCY
    places = int(payslip.get("decimal_places") or DEFAULT_DECIMAL_PLACES)
    status = "Locked" if payslip.get("is_finalized") else (payslip.get("status") or "Draft")
    deductions = payslip.get("deductions") or {}

    def esc(value: Any) -> str:
        return html_lib.escape(str(value if value is not None else ""))

    def money_text(value: Any) -> str:
        return esc(_pdf_money(value, currency, places))

    def qty_text(value: Any) -> str:
        number = float(value or 0)
        if abs(number - round(number)) < 0.001:
            return str(int(round(number)))
        return "{:.2f}".format(number)

    rows = []
    for item in _payslip_invoice_items(payslip):
        rows.append(
            "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                esc(item["description"]),
                esc(qty_text(item["qty"])),
                money_text(item["unit"]),
                money_text(item["amount"]),
            )
        )
    paid_leave = emp.get("paid_leave_days")
    if paid_leave is None:
        paid_leave = emp.get("approved_leave_days") or 0
    note = "This payslip is for the month shown above. Locked months do not change after Finalize."
    return (
        '<div class="invoice-paper" id="printPayslip"><header>'
        '<div class="invoice-brand">'
        '<img src="/static/Logo.png" alt="Company logo" onerror="this.style.display=\'none\';this.nextElementSibling.style.display=\'flex\';">'
        '<div class="payslip-mark" style="display:none">Q</div>'
        "<div><strong style=\"font-size:1.25rem\">{name}</strong>"
        '<div style="color:#247eae;font-size:.72rem;font-weight:800;letter-spacing:.5px">{tagline}</div>'
        "<div>{phone} · {email}</div>"
        '<div class="mono">NTN: {ntn}</div>'
        "<div>{address}</div><div>{website}</div></div></div>"
        '<div style="text-align:right"><h1 style="font-size:1.8rem;margin:0">PAYSLIP</h1>'
        '<div class="mono">{month}</div><span class="badge">{status}</span></div></header>'
        '<div class="invoice-meta"><div><small>Pay to</small><strong>{staff}</strong>'
        "<div>#{emp_id} · {dept} · {role}</div></div>"
        "<div><small>Pay period</small><strong>{period}</strong>"
        "<div>Work days: {work} · Days paid: {paid}</div></div></div>"
        "<table><thead><tr><th>Description</th><th>Qty</th><th>Unit price</th><th>Amount</th></tr></thead>"
        "<tbody>{rows}</tbody></table>"
        '<div class="total"><div>Gross pay: {gross}</div><div>Total cuts: {cuts}</div><hr>'
        "<strong>Take-home pay: {net}</strong></div>"
        '<div class="bank-details"><small>Payment details</small>'
        "<strong>Account Title:</strong> {bank_title}<br>"
        "<strong>Bank Name:</strong> {bank_name}<br>"
        "<strong>IBAN:</strong> {iban}</div>"
        '<p style="margin-top:22px;color:#628095">{note} Present {present} · Absent {absent} · Late {late} · '
        "Paid leave {leave} · Unpaid leave {unpaid}</p>"
        '<div class="app-preview-footer">{footer}<br><strong>{powered}</strong></div></div>'
    ).format(
        name=esc(lh["payslip_company_name"]),
        tagline=esc(lh["payslip_tagline"]),
        phone=esc(lh["payslip_phone"]),
        email=esc(lh["payslip_email"]),
        ntn=esc(lh["payslip_ntn"]),
        address=esc(lh["payslip_address"]),
        website=esc(lh["payslip_website"]),
        month=esc(payslip.get("payroll_month") or ""),
        status=esc(status),
        staff=esc(emp.get("employee_name") or "Staff"),
        emp_id=esc(emp.get("employee_id") or ""),
        dept=esc(emp.get("department") or "General"),
        role=esc(emp.get("designation") or "Staff"),
        period=esc("{} to {}".format(payslip.get("period_start") or "", payslip.get("period_end") or "")),
        work=esc(emp.get("month_scheduled_workdays") or emp.get("scheduled_workdays") or 0),
        paid=esc(emp.get("payable_days") or 0),
        rows="".join(rows),
        gross=money_text(payslip.get("gross_salary")),
        cuts=money_text(deductions.get("total")),
        net=money_text(payslip.get("net_salary")),
        bank_title=esc(lh["payslip_bank_title"]),
        bank_name=esc(lh["payslip_bank_name"]),
        iban=esc(lh["payslip_bank_iban"]),
        note=esc(note),
        present=esc(emp.get("present_days") or 0),
        absent=esc(emp.get("absent_days") or 0),
        late=esc(emp.get("late_days") or 0),
        leave=esc(paid_leave or 0),
        unpaid=esc(emp.get("unpaid_leave_days") or 0),
        footer=esc(lh["payslip_footer"]),
        powered=esc(lh["payslip_powered"]),
    )


def payslip_pdf_bytes(payslip: Dict[str, Any]) -> bytes:
    """A4 payslip PDF in the same layout as the on-screen invoice paper."""
    emp = payslip.get("employee") or {}
    lh = _payslip_letterhead(payslip)
    currency = payslip.get("currency_code") or DEFAULT_CURRENCY
    places = int(payslip.get("decimal_places") or DEFAULT_DECIMAL_PLACES)
    status = "Locked" if payslip.get("is_finalized") else (payslip.get("status") or "Draft")
    items = _payslip_invoice_items(payslip)
    deductions = payslip.get("deductions") or {}
    ops: List[str] = []

    def rgb_fill(r: float, g: float, b: float, x: float, y: float, w: float, h: float) -> None:
        ops.append("{:.3f} {:.3f} {:.3f} rg {:.1f} {:.1f} {:.1f} {:.1f} re f".format(r, g, b, x, y, w, h))

    def rgb_stroke(r: float, g: float, b: float, x: float, y: float, w: float, h: float) -> None:
        ops.append("{:.3f} {:.3f} {:.3f} RG {:.1f} {:.1f} {:.1f} {:.1f} re S".format(r, g, b, x, y, w, h))

    def text_at(x: float, y: float, text: str, size: int = 10, bold: bool = False, color: Tuple[float, float, float] = (0.09, 0.19, 0.26)) -> None:
        font = "F2" if bold else "F1"
        ops.append("{:.3f} {:.3f} {:.3f} rg".format(*color))
        ops.append("BT /{} {} Tf {:.1f} {:.1f} Td ({}) Tj ET".format(font, size, x, y, _pdf_escape(text)))

    # A4 invoice paper: header, 4-column table, totals, bank box, dark footer.
    rgb_fill(1, 1, 1, 0, 0, 595, 842)
    text_at(48, 800, lh["payslip_company_name"], 16, True, (0.09, 0.19, 0.26))
    text_at(48, 784, lh["payslip_tagline"][:72], 8, True, (0.141, 0.494, 0.682))
    text_at(48, 768, "{}  ·  {}".format(lh["payslip_phone"], lh["payslip_email"]), 9, False, (0.384, 0.502, 0.584))
    text_at(48, 754, "NTN: {}".format(lh["payslip_ntn"]), 9, True, (0.384, 0.502, 0.584))
    text_at(48, 740, lh["payslip_address"], 9, False, (0.384, 0.502, 0.584))
    text_at(430, 798, "PAYSLIP", 22, True, (0.09, 0.19, 0.26))
    text_at(430, 778, str(payslip.get("payroll_month") or ""), 11, False, (0.384, 0.502, 0.584))
    text_at(430, 760, status, 10, True, (0.239, 0.624, 0.831))
    rgb_fill(0.239, 0.624, 0.831, 36, 726, 523, 2.2)

    text_at(48, 708, "PAY TO", 8, True, (0.384, 0.502, 0.584))
    text_at(48, 692, str(emp.get("employee_name") or "Staff"), 12, True)
    text_at(48, 676, "#{}  ·  {}  ·  {}".format(
        emp.get("employee_id") or "",
        emp.get("department") or "General",
        emp.get("designation") or "Staff",
    ), 9, False, (0.384, 0.502, 0.584))
    text_at(320, 708, "PAY PERIOD", 8, True, (0.384, 0.502, 0.584))
    text_at(320, 692, "{} to {}".format(payslip.get("period_start") or "", payslip.get("period_end") or ""), 11, True)
    text_at(320, 676, "Work days: {}   Days paid: {}".format(
        emp.get("month_scheduled_workdays") or emp.get("scheduled_workdays") or 0,
        emp.get("payable_days") or 0,
    ), 9, False, (0.384, 0.502, 0.584))
    rgb_stroke(0.863, 0.925, 0.957, 36, 664, 523, 0.8)

    y = 640
    rgb_fill(0.961, 0.984, 1.0, 36, y - 6, 523, 22)
    text_at(48, y, "Description", 9, True)
    text_at(286, y, "Qty", 9, True)
    text_at(348, y, "Unit price", 9, True)
    text_at(456, y, "Amount", 9, True)
    y -= 24
    for item in items:
        text_at(48, y, str(item["description"])[:40], 10)
        qty = item["qty"]
        qty_label = str(int(qty)) if abs(qty - round(qty)) < 0.001 else "{:.2f}".format(qty)
        text_at(286, y, qty_label, 10)
        text_at(348, y, _pdf_money(item["unit"], currency, places), 10)
        text_at(448, y, _pdf_money(item["amount"], currency, places), 10, True)
        y -= 16
        if y < 250:
            break

    y = min(y - 10, 300)
    rgb_fill(0.961, 0.984, 1.0, 300, y - 70, 259, 86)
    rgb_stroke(0.812, 0.910, 0.969, 300, y - 70, 259, 86)
    text_at(314, y, "Gross pay", 9, False, (0.384, 0.502, 0.584))
    text_at(400, y, _pdf_money(payslip.get("gross_salary"), currency, places), 10, True)
    text_at(314, y - 16, "Total cuts", 9, False, (0.384, 0.502, 0.584))
    text_at(400, y - 16, _pdf_money(deductions.get("total"), currency, places), 10, True)
    text_at(314, y - 40, "Take-home pay", 11, True)
    text_at(390, y - 40, _pdf_money(payslip.get("net_salary"), currency, places), 12, True, (0.09, 0.19, 0.26))

    rgb_fill(0.961, 0.984, 1.0, 36, 86, 250, 110)
    rgb_stroke(0.812, 0.910, 0.969, 36, 86, 250, 110)
    text_at(48, 176, "PAYMENT DETAILS", 8, True, (0.384, 0.502, 0.584))
    text_at(48, 160, "Account Title: {}".format(lh["payslip_bank_title"])[:42], 9)
    text_at(48, 146, "Bank Name: {}".format(lh["payslip_bank_name"])[:42], 9)
    text_at(48, 132, "IBAN: {}".format(lh["payslip_bank_iban"])[:36], 9)
    text_at(48, 114, "Present {} · Absent {} · Late {}".format(
        emp.get("present_days") or 0,
        emp.get("absent_days") or 0,
        emp.get("late_days") or 0,
    ), 8, False, (0.384, 0.502, 0.584))

    rgb_fill(0.09, 0.19, 0.26, 36, 36, 523, 40)
    text_at(48, 58, lh["payslip_footer"][:78], 8, False, (1, 1, 1))
    text_at(48, 44, lh["payslip_powered"][:60], 9, True, (1, 1, 1))

    stream = "\n".join(ops).encode("latin-1", "replace")
    objs = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        (
            b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R /F2 6 0 R >> >> >> endobj\n"
        ),
        b"4 0 obj << /Length %d >> stream\n" % len(stream) + stream + b"\nendstream endobj\n",
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
        b"6 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >> endobj\n",
    ]
    parts = [b"%PDF-1.4\n"]
    offsets = [0]
    for obj in objs:
        offsets.append(sum(len(p) for p in parts))
        parts.append(obj)
    xref_pos = sum(len(p) for p in parts)
    xref = [b"xref\n0 7\n0000000000 65535 f \n"]
    for off in offsets[1:]:
        xref.append("{:010d} 00000 n \n".format(off).encode("ascii"))
    parts.extend(xref)
    parts.append(b"trailer << /Size 7 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref_pos)
    return b"".join(parts)


def attendance_report_pdf_bytes(report: Dict[str, Any], filter_info: Optional[Dict[str, Any]] = None, org_name: str = "Organization") -> bytes:
    """Generate a clean, multi-page A4 PDF attendance report without external libraries."""
    records = report.get('records', [])
    date_label = report.get('date_label') or (report.get('start_date') if report.get('start_date') == report.get('end_date') else f"{report.get('start_date')} to {report.get('end_date')}") or "Today"
    is_multi_day = report.get('is_multi_day', False)
    filters = filter_info or {}
    dept_filter = filters.get('department') or 'All'
    status_filter = filters.get('status') or 'All'
    shift_filter = filters.get('shift') or 'All'
    camera_filter = filters.get('camera') or 'All'
    search_filter = filters.get('search') or ''
    gen_time = datetime.now().strftime('%d %b %Y %H:%M')

    def rgb_fill(ops: List[str], r: float, g: float, b: float, x: float, y: float, w: float, h: float) -> None:
        ops.append("{:.3f} {:.3f} {:.3f} rg {:.1f} {:.1f} {:.1f} {:.1f} re f".format(r, g, b, x, y, w, h))

    def rgb_stroke(ops: List[str], r: float, g: float, b: float, x: float, y: float, w: float, h: float) -> None:
        ops.append("{:.3f} {:.3f} {:.3f} RG {:.1f} {:.1f} {:.1f} {:.1f} re S".format(r, g, b, x, y, w, h))

    def text_at(ops: List[str], x: float, y: float, text: str, size: int = 9, bold: bool = False, color: Tuple[float, float, float] = (0.09, 0.19, 0.26)) -> None:
        font = "F2" if bold else "F1"
        ops.append("{:.3f} {:.3f} {:.3f} rg".format(*color))
        ops.append("BT /{} {} Tf {:.1f} {:.1f} Td ({}) Tj ET".format(font, size, x, y, _pdf_escape(text)))

    rows_per_page = 26
    total_records = len(records)
    total_pages = max(1, (total_records + rows_per_page - 1) // rows_per_page) if total_records > 0 else 1

    pages_content = []
    for page_num in range(1, total_pages + 1):
        ops: List[str] = []
        rgb_fill(ops, 1, 1, 1, 0, 0, 595, 842)

        # Header banner (dark slate)
        rgb_fill(ops, 0.09, 0.19, 0.26, 36, 786, 523, 36)
        text_at(ops, 48, 804, str(org_name)[:45], 13, True, (1, 1, 1))
        text_at(ops, 48, 792, "ATTENDANCE REPORT", 8, True, (0.24, 0.62, 0.83))
        text_at(ops, 380, 804, f"Period: {date_label}"[:30], 8, True, (1, 1, 1))
        text_at(ops, 380, 792, f"Generated: {gen_time}", 8, False, (0.7, 0.8, 0.85))

        # Filters subheader box
        rgb_fill(ops, 0.96, 0.98, 1.0, 36, 746, 523, 34)
        rgb_stroke(ops, 0.86, 0.92, 0.96, 36, 746, 523, 34)
        f_text1 = f"Department: {dept_filter}   |   Status: {status_filter}   |   Shift: {shift_filter}"
        f_text2 = f"Camera: {camera_filter}" + (f"   |   Search: {search_filter}" if search_filter else "") + f"   |   Total Staff: {total_records}"
        text_at(ops, 46, 764, f_text1[:90], 8, False, (0.38, 0.50, 0.58))
        text_at(ops, 46, 752, f_text2[:90], 8, False, (0.38, 0.50, 0.58))

        # Table header
        table_top = 718
        row_height = 22
        rgb_fill(ops, 0.14, 0.49, 0.68, 36, table_top, 523, row_height)
        if is_multi_day:
            text_at(ops, 42, table_top + 6, "ID", 8, True, (1, 1, 1))
            text_at(ops, 80, table_top + 6, "Employee Name", 8, True, (1, 1, 1))
            text_at(ops, 220, table_top + 6, "Department", 8, True, (1, 1, 1))
            text_at(ops, 310, table_top + 6, "Shift", 8, True, (1, 1, 1))
            text_at(ops, 390, table_top + 6, "Present / Work", 8, True, (1, 1, 1))
            text_at(ops, 465, table_top + 6, "Absent", 8, True, (1, 1, 1))
            text_at(ops, 510, table_top + 6, "Late", 8, True, (1, 1, 1))
        else:
            text_at(ops, 40, table_top + 6, "ID", 8, True, (1, 1, 1))
            text_at(ops, 76, table_top + 6, "Employee Name", 8, True, (1, 1, 1))
            text_at(ops, 190, table_top + 6, "Department", 8, True, (1, 1, 1))
            text_at(ops, 265, table_top + 6, "Status", 8, True, (1, 1, 1))
            text_at(ops, 340, table_top + 6, "Check-In", 8, True, (1, 1, 1))
            text_at(ops, 390, table_top + 6, "In Camera", 8, True, (1, 1, 1))
            text_at(ops, 455, table_top + 6, "Check-Out", 8, True, (1, 1, 1))
            text_at(ops, 505, table_top + 6, "Out Cam", 8, True, (1, 1, 1))

        # Rows on this page
        start_idx = (page_num - 1) * rows_per_page
        end_idx = min(start_idx + rows_per_page, total_records)
        page_records = records[start_idx:end_idx]

        y = table_top - row_height
        for idx, rec in enumerate(page_records):
            bg = (0.98, 0.99, 1.0) if idx % 2 == 1 else (1.0, 1.0, 1.0)
            rgb_fill(ops, bg[0], bg[1], bg[2], 36, y, 523, row_height)
            rgb_stroke(ops, 0.90, 0.93, 0.96, 36, y, 523, row_height)

            emp_id = str(rec.get('user_id') or '')
            name = str(rec.get('name') or '')
            dept = str(rec.get('department') or 'General')

            if is_multi_day:
                shift_name = str(rec.get('shift_name') or 'General')
                pres_disp = str(rec.get('present_display') or f"{rec.get('present_days', 0)} / {rec.get('working_days', 0)}")
                abs_disp = str(rec.get('absent_display') or str(rec.get('absent_days', 0)))
                late_disp = str(rec.get('late_days') or 0)

                text_at(ops, 42, y + 6, f"#{emp_id}", 8, True, (0.09, 0.19, 0.26))
                text_at(ops, 80, y + 6, name[:24], 8, False, (0.09, 0.19, 0.26))
                text_at(ops, 220, y + 6, dept[:15], 8, False, (0.38, 0.50, 0.58))
                text_at(ops, 310, y + 6, shift_name[:14], 8, False, (0.38, 0.50, 0.58))
                text_at(ops, 390, y + 6, pres_disp[:14], 8, True, (0.12, 0.55, 0.24))
                text_at(ops, 465, y + 6, abs_disp[:8], 8, False, (0.85, 0.19, 0.15))
                text_at(ops, 510, y + 6, late_disp[:6], 8, False, (0.69, 0.38, 0.0))
            else:
                st = str(rec.get('status_display') or rec.get('status') or 'Absent')
                first_seen = str(rec.get('first_seen') or '-')
                if ' ' in first_seen:
                    first_seen = first_seen.split(' ')[-1][:5]
                last_seen = str(rec.get('last_seen') or '-')
                if ' ' in last_seen:
                    last_seen = last_seen.split(' ')[-1][:5]
                in_cam = str(rec.get('checkin_camera') or rec.get('location') or '-')
                out_cam = str(rec.get('checkout_camera') or rec.get('checkout_location') or '-')

                st_color = (0.12, 0.55, 0.24) if 'present' in st.lower() else ((0.69, 0.38, 0.0) if 'late' in st.lower() else (0.85, 0.19, 0.15))

                text_at(ops, 40, y + 6, f"#{emp_id}", 8, True, (0.09, 0.19, 0.26))
                text_at(ops, 76, y + 6, name[:20], 8, False, (0.09, 0.19, 0.26))
                text_at(ops, 190, y + 6, dept[:12], 8, False, (0.38, 0.50, 0.58))
                text_at(ops, 265, y + 6, st[:15], 8, True, st_color)
                text_at(ops, 340, y + 6, first_seen[:8], 8, False, (0.09, 0.19, 0.26))
                text_at(ops, 390, y + 6, in_cam[:12], 7, False, (0.38, 0.50, 0.58))
                text_at(ops, 455, y + 6, last_seen[:8], 8, False, (0.09, 0.19, 0.26))
                text_at(ops, 505, y + 6, out_cam[:10], 7, False, (0.38, 0.50, 0.58))

            y -= row_height

        if total_records == 0:
            rgb_fill(ops, 1, 1, 1, 36, table_top - row_height, 523, 40)
            text_at(ops, 160, table_top - 20, "No attendance records found matching the active filters.", 9, False, (0.5, 0.5, 0.5))

        # Footer
        rgb_fill(ops, 0.09, 0.19, 0.26, 36, 30, 523, 20)
        text_at(ops, 48, 36, "QuantEye Security & Attendance System", 8, False, (1, 1, 1))
        text_at(ops, 480, 36, f"Page {page_num} of {total_pages}", 8, True, (1, 1, 1))

        pages_content.append("\n".join(ops))

    font1_id = 3 + len(pages_content) * 2
    font2_id = font1_id + 1
    total_objs = font2_id

    kids = []
    page_objs = []
    for i, content in enumerate(pages_content):
        p_id = 3 + i * 2
        c_id = p_id + 1
        kids.append(f"{p_id} 0 R")
        stream = content.encode("latin-1", "replace")
        p_obj = (
            f"{p_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Contents {c_id} 0 R /Resources << /Font << /F1 {font1_id} 0 R /F2 {font2_id} 0 R >> >> >> endobj\n"
        ).encode("latin-1")
        c_obj = f"{c_id} 0 obj << /Length {len(stream)} >> stream\n".encode("latin-1") + stream + b"\nendstream endobj\n"
        page_objs.extend([p_obj, c_obj])

    kids_str = " ".join(kids)
    objs = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{kids_str}] /Count {len(pages_content)} >> endobj\n".encode("latin-1"),
        *page_objs,
        b"%d 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n" % font1_id,
        b"%d 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >> endobj\n" % font2_id,
    ]
    parts = [b"%PDF-1.4\n"]
    offsets = [0]
    for obj in objs:
        offsets.append(sum(len(p) for p in parts))
        parts.append(obj)
    xref_pos = sum(len(p) for p in parts)
    xref = [f"xref\n0 {total_objs + 1}\n0000000000 65535 f \n".encode("ascii")]
    for off in offsets[1:]:
        xref.append("{:010d} 00000 n \n".format(off).encode("ascii"))
    parts.extend(xref)
    parts.append(f"trailer << /Size {total_objs + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%%%EOF\n".encode("ascii"))
    return b"".join(parts)
